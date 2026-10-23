from collections.abc import Callable
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from django import forms
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from country_workspace.datasources.rdi.config import Config, SheetName
from country_workspace.datasources.rdi.processors import read_sheets
from country_workspace.utils import import_template as tpl
from country_workspace.workspaces.admin.forms import ImportFileForm

Checker = SimpleNamespace
Factory = Callable[..., Checker]


@pytest.fixture
def checker() -> Factory:
    def _make(**fields: forms.Field) -> Checker:
        form_class = type("StubForm", (forms.Form,), dict(fields))
        return SimpleNamespace(get_form_class=lambda: form_class)

    return _make


@pytest.fixture
def program() -> Callable[..., Any]:
    def _make(*, master_detail: bool | None, household: Checker | None, individual: Checker | None) -> Any:
        group = None if master_detail is None else SimpleNamespace(master_detail=master_detail)
        return SimpleNamespace(
            beneficiary_group=group,
            household_checker=household,
            individual_checker=individual,
            country_office=None,
        )

    return _make


@pytest.fixture
def household_checker(checker: Factory) -> Checker:
    return checker(
        household_id=forms.CharField(),
        size=forms.IntegerField(min_value=1, max_value=50),
        head_of_household_id=forms.IntegerField(required=False),
        consent=forms.BooleanField(required=False),
        children_count=forms.IntegerField(required=False),
    )


@pytest.fixture
def individual_checker(checker: Factory) -> Checker:
    return checker(
        individual_id=forms.CharField(),
        household_id=forms.CharField(),
        full_name=forms.CharField(label="Full name", help_text="As written on the ID"),
        birth_date=forms.DateField(),
        gender=forms.ChoiceField(choices=[("MALE", "Male"), ("FEMALE", "Female")]),
        flags=forms.MultipleChoiceField(choices=[("A", "a"), ("B", "b")], required=False),
        weight=forms.DecimalField(required=False),
        phone_no=forms.CharField(required=False),
    )


def build(program: Any) -> Workbook:
    return load_workbook(tpl.build_import_template(program))


def headers(sheet: Worksheet) -> list[str]:
    return [cell.value for cell in sheet[1]]


def validations(sheet: Worksheet) -> dict[str, Any]:
    return {str(dv.sqref)[0]: dv for dv in sheet.data_validations.dataValidation}


@pytest.mark.parametrize(
    ("master_detail", "household", "individual", "expected"),
    [
        (True, True, True, [SheetName.HOUSEHOLDS, SheetName.INDIVIDUALS]),
        (True, True, False, [SheetName.HOUSEHOLDS]),
        (True, False, True, [SheetName.INDIVIDUALS]),
        (True, False, False, []),
        (False, True, True, [SheetName.PEOPLE]),
        (False, True, False, []),
        (None, True, True, [SheetName.PEOPLE]),
    ],
    ids=["md-both", "md-household-only", "md-individual-only", "md-none", "people", "people-no-individual", "no-group"],
)
def test_sheets_follow_the_selected_data_checkers(
    program: Callable[..., Any],
    household_checker: Checker,
    individual_checker: Checker,
    master_detail: bool | None,
    household: bool,
    individual: bool,
    expected: list[SheetName],
) -> None:
    prog = program(
        master_detail=master_detail,
        household=household_checker if household else None,
        individual=individual_checker if individual else None,
    )

    assert [plan.name for plan in tpl.get_sheet_plans(prog)] == expected


def test_build_requires_a_data_checker(program: Callable[..., Any]) -> None:
    with pytest.raises(tpl.NoDataCheckerError, match="no data checker"):
        tpl.build_import_template(program(master_detail=True, household=None, individual=None))


def test_master_detail_has_two_visible_pages_and_hidden_choices(
    program: Callable[..., Any], household_checker: Checker, individual_checker: Checker
) -> None:
    wb = build(program(master_detail=True, household=household_checker, individual=individual_checker))

    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [
        ("Households", "visible"),
        ("Individuals", "visible"),
        ("Choices", "hidden"),
    ]


def test_people_only_has_one_visible_page_with_prefixed_columns(
    program: Callable[..., Any], individual_checker: Checker
) -> None:
    wb = build(program(master_detail=False, household=None, individual=individual_checker))

    assert [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] == ["People"]
    assert headers(wb["People"]) == [
        "pp_index_id",
        "pp_individual_id",
        "pp_household_id",
        "pp_full_name",
        "pp_birth_date",
        "pp_gender",
        "pp_flags",
        "pp_weight",
        "pp_phone_no",
    ]


def test_no_choices_sheet_without_choice_fields(program: Callable[..., Any], household_checker: Checker) -> None:
    wb = build(program(master_detail=True, household=household_checker, individual=None))

    assert wb.sheetnames == ["Households"]


def test_identifiers_come_first_and_are_not_duplicated(
    program: Callable[..., Any], household_checker: Checker, individual_checker: Checker
) -> None:
    wb = build(program(master_detail=True, household=household_checker, individual=individual_checker))

    assert headers(wb["Households"]) == [
        "household_id",
        "size",
        "head_of_household_id",
        "consent",
        "children_count",
    ]
    assert headers(wb["Individuals"])[:3] == ["individual_id", "household_id", "full_name"]
    assert headers(wb["Individuals"]).count("household_id") == 1


def test_header_notes_describe_the_field(
    program: Callable[..., Any], household_checker: Checker, individual_checker: Checker
) -> None:
    wb = build(program(master_detail=True, household=household_checker, individual=individual_checker))
    sheet = wb["Individuals"]

    note = sheet["C1"].comment.text
    assert "Full name" in note
    assert "As written on the ID" in note
    assert "0 for an individual without a household" in sheet["B1"].comment.text


def test_head_of_household_points_at_the_individuals_sheet(
    program: Callable[..., Any], household_checker: Checker
) -> None:
    sheet = build(program(master_detail=True, household=household_checker, individual=None))["Households"]

    assert "Individuals" in sheet["C1"].comment.text


def test_constraints_match_field_types(
    program: Callable[..., Any], household_checker: Checker, individual_checker: Checker
) -> None:
    wb = build(program(master_detail=True, household=household_checker, individual=individual_checker))

    hh = validations(wb["Households"])
    assert (hh["B"].type, hh["B"].operator, hh["B"].formula1, hh["B"].formula2) == ("whole", None, "1", "50")
    assert hh["E"].type == "whole"  # no bounds on the field, still has to be a number
    assert hh["D"].type == "list"
    assert set(hh["D"].formula1.strip('"').split(",")) == {"True", "False"}

    ind = validations(wb["Individuals"])
    assert ind["D"].type == "date"
    assert ind["G"].type == "decimal"


def test_single_choice_is_a_dropdown_and_multiple_choice_is_not(
    program: Callable[..., Any], individual_checker: Checker
) -> None:
    wb = build(program(master_detail=True, household=None, individual=individual_checker))

    ind = validations(wb["Individuals"])
    assert ind["E"].type == "list"
    assert ind["E"].formula1.startswith("Choices!")
    assert "F" not in ind
    assert [c.value for c in wb["Choices"]["A"]][1:3] == ["MALE", "FEMALE"]
    assert "A, B" in wb["Individuals"]["F1"].comment.text


def test_long_choice_lists_are_stored_on_the_choices_sheet(program: Callable[..., Any], checker: Factory) -> None:
    choices = [(f"AF{i:04d}", f"AF{i:04d} - A rather long area name {i}") for i in range(500)]
    individuals = checker(admin1=forms.ChoiceField(choices=[("", "None"), *choices]))

    wb = build(program(master_detail=True, household=None, individual=individuals))

    dv = validations(wb["Individuals"])["C"]
    assert dv.formula1 == "Choices!$A$2:$A$501"
    assert wb["Choices"].max_row == len(choices) + 1
    assert wb["Choices"]["A2"].value == "AF0000"
    assert wb["Choices"]["B2"].value == "AF0000 - A rather long area name 0"


def test_option_groups_are_flattened(program: Callable[..., Any], checker: Factory) -> None:
    individuals = checker(
        kind=forms.ChoiceField(choices=[("Group", [("A", "a"), ("B", "b")]), ("C", "c")]),
    )

    wb = build(program(master_detail=True, household=None, individual=individuals))

    assert [c.value for c in wb["Choices"]["A"]][1:] == ["A", "B", "C"]


def test_text_columns_are_formatted_as_text(program: Callable[..., Any], individual_checker: Checker) -> None:
    sheet = build(program(master_detail=True, household=None, individual=individual_checker))["Individuals"]

    assert sheet.column_dimensions["H"].number_format == "@"
    assert sheet.column_dimensions["A"].number_format == "General"


def test_header_is_frozen_and_required_columns_are_highlighted(
    program: Callable[..., Any], individual_checker: Checker
) -> None:
    sheet = build(program(master_detail=True, household=None, individual=individual_checker))["Individuals"]

    assert sheet.freeze_panes == "A2"
    assert sheet["C1"].fill.fgColor.rgb != sheet["H1"].fill.fgColor.rgb


def test_references_to_other_sheets_accept_any_identifier(
    program: Callable[..., Any], household_checker: Checker
) -> None:
    sheet = build(program(master_detail=True, household=household_checker, individual=None))["Households"]

    assert "C" not in validations(sheet)  # head_of_household_id is an integer field, but ids may be text
    assert sheet.column_dimensions["C"].number_format == "General"


def test_header_text_is_never_interpreted_as_a_formula(program: Callable[..., Any], checker: Factory) -> None:
    individuals = checker(**{"=SUM(A1)": forms.CharField(required=False)})

    sheet = build(program(master_detail=True, household=None, individual=individuals))["Individuals"]

    assert sheet["C1"].value == "=SUM(A1)"
    assert sheet["C1"].data_type == "s"


@pytest.mark.parametrize("master_detail", [True, False], ids=["master_detail", "people"])
def test_default_import_form_values_match_the_template(master_detail: bool) -> None:
    form = ImportFileForm(beneficiary_group=SimpleNamespace(master_detail=master_detail))

    assert form.fields["beneficiary_id_column"].initial == (
        tpl.INDIVIDUAL_ID_COLUMN if master_detail else tpl.PEOPLE_ID_COLUMN
    )
    if master_detail:
        assert form.fields["household_id_column"].initial == tpl.HOUSEHOLD_ID_COLUMN
        assert form.fields["household_label"].initial == tpl.HOUSEHOLD_ID_COLUMN
    else:
        assert form.fields["people_prefix"].initial == tpl.PEOPLE_PREFIX
    assert form.fields["first_line"].initial == 2


def test_filled_master_detail_template_is_read_by_the_importer(
    program: Callable[..., Any], household_checker: Checker, individual_checker: Checker, tmp_path: Any
) -> None:
    template = tpl.build_import_template(
        program(master_detail=True, household=household_checker, individual=individual_checker)
    )
    wb = load_workbook(template)
    wb["Households"].append([1, 3, 1, True])
    wb["Individuals"].append([1, 1, "John Doe", None, "MALE"])
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    config: Config = {
        "master_detail": True,
        "household_id_column": tpl.HOUSEHOLD_ID_COLUMN,
        "beneficiary_id_column": tpl.INDIVIDUAL_ID_COLUMN,
        "first_line": 2,
        "batch_name": "b",
    }

    households, individuals = (list(s) for s in read_sheets(config, str(path), "Households", "Individuals"))

    assert households == [
        {"household_id": 1, "size": 3, "head_of_household_id": 1, "consent": True, "children_count": None}
    ]
    assert [i["full_name"] for i in individuals] == ["John Doe"]
    assert individuals[0]["household_id"] == households[0]["household_id"]


def test_filled_people_template_is_read_by_the_importer(
    program: Callable[..., Any], individual_checker: Checker, tmp_path: Any
) -> None:
    template = tpl.build_import_template(program(master_detail=False, household=None, individual=individual_checker))
    wb = load_workbook(BytesIO(template.getvalue()))
    wb["People"].append([1, "p-1", None, "Jane Doe"])
    path = tmp_path / "people.xlsx"
    wb.save(path)
    config: Config = {
        "master_detail": False,
        "beneficiary_id_column": tpl.PEOPLE_ID_COLUMN,
        "people_prefix": tpl.PEOPLE_PREFIX,
        "first_line": 2,
        "batch_name": "b",
    }

    (people,) = (list(s) for s in read_sheets(config, str(path), "People"))

    assert people[0]["pp_index_id"] == 1
    assert people[0]["pp_full_name"] == "Jane Doe"
